#!/usr/bin/env python3
"""
banco.py: gera o banco de dengue do épico 2 (Banco de Dados, A59 a A63).

Lê dengue_pr_limpo.xlsx (congelado com os números baixados em 09/09/2026; este
script NÃO acessa a rede) e escreve, em scripts/dataset-dengue/servidor/:

  initdb/01-dengue.sql   carga do Postgres: 32 contas de aluno, 1 banco modelo
                         com as duas tabelas, e 32 bancos dengue_NN copiados
                         dele, cada um com o aluno como dono
  esperado.env           totais de conferência, lidos pelo verificar.sh

Decisão de 22/09/2026: o épico roda no Postgres da máquina do professor, com os
alunos entrando pelo DBeaver (mesmo esquema da A47), e não no sqliteonline.
O mesmo servidor atende os Indicadores 5 (consultas), 4 (CREATE ROLE, GRANT,
REVOKE) e 6 (pg_dump e restore).

Um banco POR ALUNO, não por dupla (decisão de 23/09): no banco compartilhado da
A47, um mexia e o outro ficava travado sem saber por quê. Aluno enrolado se
resolve com `servidor/resetar.sh NN`, que refaz só o banco dele.

Duas tabelas:
  casos_dengue  408 linhas, uma por município, ano e mês. Tem municipio e
                macrorregional junto, igual à planilha, para a A59 fazer WHERE
                e GROUP BY sem precisar de JOIN.
  municipios    17 linhas. A população mora só aqui: casos por 100 mil
                habitantes exige JOIN, e esse é o gancho da A60.

Uso:  venv/bin/python scripts/dataset-dengue/banco.py
Depois de gerar de novo, o servidor só recarrega com o volume vazio:
      cd scripts/dataset-dengue/servidor && docker compose down -v && docker compose up -d
"""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[2]
ENTRADA = RAIZ / "aulas/09set/A55_UC01_11set/public/dados/dengue_pr_limpo.xlsx"
SERVIDOR = Path(__file__).resolve().parent / "servidor"
SAIDA_SQL = SERVIDOR / "initdb/01-dengue.sql"
SAIDA_ENV = SERVIDOR / "esperado.env"

ALUNOS = 32  # 30 na chamada + folga. NN = número da chamada
REGIOES = ("Leste", "Noroeste", "Norte", "Oeste")
ANOS = (2024, 2025)


def ler_limpo():
    ws = load_workbook(ENTRADA, read_only=True)["01_dados"]
    linhas = list(ws.iter_rows(values_only=True))
    cabecalho, corpo = linhas[0], linhas[1:]
    registros = [dict(zip(cabecalho, l)) for l in corpo if l[0] is not None]
    if len(registros) != 408:
        sys.exit(f"ERRO: esperava 408 linhas limpas, achei {len(registros)}")
    return registros


def municipios_de(registros):
    """Um por código IBGE. Para se o mesmo município tiver dois nomes, duas
    regiões ou duas populações, porque aí a tabela municipios mentiria."""
    vistos = {}
    for r in registros:
        chave = r["codigo_ibge"]
        m = (r["municipio"], r["macrorregional"], r["populacao"])
        if vistos.setdefault(chave, m) != m:
            sys.exit(f"ERRO: município {chave} aparece com dados diferentes: {vistos[chave]} e {m}")
        if r["macrorregional"] not in REGIOES:
            sys.exit(f"ERRO: macrorregional fora das 4: {r['macrorregional']!r}")
    if len(vistos) != 17:
        sys.exit(f"ERRO: esperava 17 municípios, achei {len(vistos)}")
    return dict(sorted(vistos.items(), key=lambda kv: kv[1][0]))


def texto(s):
    return "'" + str(s).replace("'", "''") + "'"


def gerar_sql(registros, municipios):
    nn = [f"{i:02d}" for i in range(1, ALUNOS + 1)]
    out = []
    w = out.append

    w("-- ---------------------------------------------------------------------------")
    w("-- Banco de dengue do épico 2 (Banco de Dados, A59 a A63)")
    w("-- GERADO por scripts/dataset-dengue/banco.py. Não edite à mão: rode o script.")
    w("--")
    w("-- Cada aluno entra com usuario alunoNN e senha dengueNN, no banco dengue_NN")
    w("-- (NN = número da chamada). O aluno é DONO do próprio banco e tem CREATEROLE e")
    w("-- CREATEDB: consegue criar contas, dar GRANT e REVOKE (Indicador 4) e restaurar")
    w("-- backup num banco novo (Indicador 6). Sem SUPERUSER: não mexe no banco dos outros.")
    w("-- ---------------------------------------------------------------------------")
    w("")
    w("-- ---------- 1. Contas dos alunos ----------")
    for n in nn:
        w(f"CREATE ROLE aluno{n} LOGIN PASSWORD 'dengue{n}' CREATEROLE CREATEDB;")
    w("")
    w("-- ---------- 2. Banco modelo ----------")
    w("CREATE DATABASE dengue_modelo;")
    w("\\connect dengue_modelo")
    w("")
    w("CREATE TABLE municipios (")
    w("    codigo_ibge     INTEGER PRIMARY KEY,")
    w("    municipio       TEXT    NOT NULL UNIQUE,")
    w(f"    macrorregional  TEXT    NOT NULL CHECK (macrorregional IN ({', '.join(texto(r) for r in REGIOES)})),")
    w("    populacao       INTEGER NOT NULL CHECK (populacao > 0)")
    w(");")
    w("")
    w("CREATE TABLE casos_dengue (")
    w("    codigo_ibge      INTEGER NOT NULL REFERENCES municipios (codigo_ibge),")
    w("    municipio        TEXT    NOT NULL,")
    w("    macrorregional   TEXT    NOT NULL,")
    w("    ano              INTEGER NOT NULL,")
    w("    mes              INTEGER NOT NULL CHECK (mes BETWEEN 1 AND 12),")
    w("    data_referencia  DATE    NOT NULL,")
    w("    casos            INTEGER NOT NULL CHECK (casos >= 0),")
    w("    PRIMARY KEY (codigo_ibge, ano, mes)")
    w(");")
    w("")
    w("INSERT INTO municipios (codigo_ibge, municipio, macrorregional, populacao) VALUES")
    w(",\n".join(f" ({c}, {texto(m)}, {texto(r)}, {p})" for c, (m, r, p) in municipios.items()) + ";")
    w("")
    w("INSERT INTO casos_dengue (codigo_ibge, municipio, macrorregional, ano, mes, data_referencia, casos) VALUES")
    ordenados = sorted(registros, key=lambda r: (r["municipio"], r["ano"], r["mes"]))
    w(",\n".join(
        f" ({r['codigo_ibge']}, {texto(r['municipio'])}, {texto(r['macrorregional'])}, "
        f"{r['ano']}, {r['mes']}, '{r['data_referencia']:%Y-%m-%d}', {r['casos']})"
        for r in ordenados) + ";")
    w("")
    w("\\connect postgres")
    w("")
    w("-- ---------- 3. Um banco por aluno, copiado do modelo ----------")
    for n in nn:
        w(f"CREATE DATABASE dengue_{n} TEMPLATE dengue_modelo OWNER aluno{n};")
    w("")
    w("-- ---------- 4. Cada aluno vira dono das tabelas do próprio banco ----------")
    w("-- Sem isso as tabelas continuam do superusuário e o aluno não consegue dar GRANT.")
    w("-- (o resetar.sh repete este passo para um aluno só)")
    for n in nn:
        w(f"\\connect dengue_{n}")
        w(f"ALTER TABLE municipios OWNER TO aluno{n}; ALTER TABLE casos_dengue OWNER TO aluno{n};")
    w("")
    w("\\connect postgres")
    w(f"\\echo '--- carga concluida: {ALUNOS} alunos, {ALUNOS} bancos dengue_NN ---'")
    return "\n".join(out) + "\n"


def gerar_env(registros):
    total = defaultdict(int)
    for r in registros:
        total[r["ano"]] += r["casos"]
    londrina_2025 = sum(r["casos"] for r in registros
                        if r["municipio"] == "Londrina" and r["ano"] == 2025)
    return (
        "# GERADO por scripts/dataset-dengue/banco.py. Conferido pelo verificar.sh.\n"
        f"ALUNOS={ALUNOS}\n"
        f"LINHAS=408\n"
        f"MUNICIPIOS=17\n"
        f"TOTAL_2024={total[2024]}\n"
        f"TOTAL_2025={total[2025]}\n"
        f"LONDRINA_2025={londrina_2025}\n"
    )


def main():
    registros = ler_limpo()
    municipios = municipios_de(registros)
    SAIDA_SQL.parent.mkdir(parents=True, exist_ok=True)
    SAIDA_SQL.write_text(gerar_sql(registros, municipios), encoding="utf-8")
    SAIDA_ENV.write_text(gerar_env(registros), encoding="utf-8")
    print(f"ok  {SAIDA_SQL.relative_to(RAIZ)}")
    print(f"ok  {SAIDA_ENV.relative_to(RAIZ)}")
    print(SAIDA_ENV.read_text(), end="")


if __name__ == "__main__":
    main()
