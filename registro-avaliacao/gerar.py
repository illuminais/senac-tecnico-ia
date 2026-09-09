#!/usr/bin/env python3
"""
Gera Office Scripts para lançar as menções do 2º trimestre no arquivo oficial.

    python gerar.py frases            # cria/atualiza o rascunho de frases.json
    python gerar.py UC2 --dry-run     # relatório do que seria escrito (não escreve nada)
    python gerar.py UC2               # emite saida/preencher-UC2.ts

O Office Script não lê arquivo local, então as menções vão embutidas no .ts gerado.
Rode de novo a qualquer momento: os indicadores ainda sem menção no CSV são ignorados.
"""
import argparse
import csv
import json
import os
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

BASE = Path(__file__).resolve().parent
CSV_PATH = Path(os.environ.get("T2_CSV", Path.home() / "Downloads" / "TECIA-2TRI-indicadores.csv"))
XLSX_PATH = Path(os.environ.get("T2_XLSX", Path.home() / "Downloads" / "Inteligencia Artificial_2026_ 1º_Ano.xlsx"))
FRASES_JSON = BASE / "frases.json"
MODELO_TS = BASE / "modelo.ts"
SAIDA = BASE / "saida"

# A numeração das abas não é a da planilha do instrutor. Verificado pelo campo CC/UC de cada aba.
TAB2UC = {
    "UC1": ("UC04", "Fundamentos e Conceitos de IA"),
    "UC2": ("UC05", "Desenvolvimento de Linguagem Python"),
    "UC3": ("UC06", "Arquitetura de Computadores e GPU"),
    "UC4": ("UC07", "Transformação Digital"),
    "UC5": ("UC08", "Desenvolvimento de Banco de Dados"),
    "UC6": ("UC09", "Estatística Aplicada e Lógica Matemática"),
    "UC7": ("UC01", "Fundamentos de Computação"),
    "UC8": ("UC02", "Inglês Instrumental"),
    "UC9": ("UC03", "Fundamentos Matemáticos para Computação e IA"),
}

# grafias divergentes entre o CSV e a planilha; a chave é o nome normalizado do CSV
ALIAS = {
    "EVADIDA DA SILVA": "JULLYA EDUARDA ALBANI DA SILVA",
    "VITORIA HENNIG BESSA": "VITÓRIA HENNIH BESSA",
}

# evadida: NAV em tudo, nenhuma linha no Feedback de Processo
EVADIDA = "JULLYA EDUARDA ALBANI DA SILVA"

MENCOES_VALIDAS = ("A", "PA", "NA", "NAV")

NAV_TXT = "Não houve evidências suficientes para avaliação deste indicador."
PROC_PA = "Demonstrou domínio parcial dos critérios. Foi ofertada recuperação, não realizada pelo(a) estudante."
PROC_NA = "Não atingiu aproveitamento suficiente. Foi ofertada recuperação, não realizada pelo(a) estudante."

# verbos em 3ª pessoa usados na abertura das descrições de indicador
VERBOS = {
    "acessa", "age", "analisa", "aplica", "armazena", "avalia", "classifica", "compreende",
    "considera", "cria", "demonstra", "diferencia", "edita", "elabora", "emprega",
    "estuda", "explora", "formata", "gerencia", "identifica", "imprime", "interpreta",
    "manipula", "organiza", "propõe", "realiza", "reconhece", "relaciona", "representa",
    "resolve",
    "seleciona",
    "usa", "utiliza", "verifica",
}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().upper()


def limpa(s):
    """desfaz o escape do export do Excel e normaliza espaços"""
    s = str(s or "").replace("_x0002_", "-").replace("_x0000_", "")
    return re.sub(r"\s+", " ", s).strip()


# ---------------------------------------------------------------------------
# flexão da descrição do indicador conforme a menção
# ---------------------------------------------------------------------------

def grupo_verbal(texto):
    """quantos tokens iniciais formam o grupo verbal ('Compreende e aplica' -> 3)"""
    toks = texto.split()
    fim = 0
    i = 0
    while i < len(toks):
        t = toks[i].strip(",").lower()
        if t in VERBOS:
            fim = i + 1
            i += 1
        elif t in ("e", "ou") and i + 1 < len(toks) and toks[i + 1].strip(",").lower() in VERBOS:
            i += 1
        else:
            break
    return fim


def flexionar(texto):
    """gera as 4 variantes a partir da descrição do indicador (rascunho, revisar à mão)"""
    texto = limpa(texto)
    if not texto.endswith("."):
        texto += "."
    toks = texto.split()
    n = grupo_verbal(texto)
    if n == 0:  # não reconheceu o verbo: devolve marcado para revisão
        return {"A": texto, "PA": "?? " + texto, "NA": "?? " + texto, "NAV": NAV_TXT}

    cabeca, resto = toks[:n], toks[n:]

    pa = " ".join(cabeca + ["parcialmente"] + resto)

    na_toks = []
    for i, t in enumerate(cabeca):
        base = t.strip(",")
        virgula = "," if t.endswith(",") else ""
        if base.lower() in VERBOS:
            palavra = base.lower()
            na_toks.append(("Não " if i == 0 else "não ") + palavra + virgula)
        else:
            na_toks.append(t)
    na = " ".join(na_toks + ["adequadamente"] + resto)

    return {"A": texto, "PA": pa, "NA": na, "NAV": NAV_TXT}


# ---------------------------------------------------------------------------
# leitura das fontes
# ---------------------------------------------------------------------------

def ler_csv():
    """-> (ordem de nomes na grafia da planilha, {nome: {(UCxx, ind): menção}})"""
    with open(CSV_PATH, encoding="utf-8") as fh:
        linhas = list(csv.reader(fh))
    ucs, nums = linhas[0], linhas[3]
    colunas = {}
    for c in range(2, min(31, len(ucs))):
        if ucs[c].strip() and nums[c].strip():
            colunas[c] = (ucs[c].strip(), nums[c].strip())

    ordem, mencoes = [], {}
    for linha in linhas[6:]:
        if len(linha) < 2 or not linha[1].strip():
            continue
        bruto = linha[1].strip()
        nome = ALIAS.get(norm(bruto), bruto)
        ordem.append(nome)
        m = {}
        for c, (uc, num) in colunas.items():
            v = linha[c].strip().upper() if c < len(linha) else ""
            if v:
                m[(uc, num)] = v
        mencoes[nome] = m
    return ordem, mencoes


def abrir_wb():
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl não instalado. Rode:  ./venv/bin/python gerar.py ...")
    return openpyxl.load_workbook(XLSX_PATH, data_only=True)


def abas_da_uc(wb, tab, trimestre=2):
    """-> (registro do trimestre, feedback de indicador, feedback de processo) com o nome exato"""
    achadas = {}
    for nome in wb.sheetnames:
        prefixo = nome.split("-")[0].strip()
        if prefixo != tab:
            continue
        if f"Registro de Avaliação {trimestre}" in nome:
            achadas["reg"] = nome
        elif "Feedback de Indicador" in nome:
            achadas["ind"] = nome
        elif "Feedback de Processo" in nome:
            achadas["proc"] = nome
    faltando = [k for k in ("reg", "ind", "proc") if k not in achadas]
    if faltando:
        sys.exit(f"aba(s) não encontrada(s) para {tab} no {trimestre}º trimestre: {faltando}")
    return achadas["reg"], achadas["ind"], achadas["proc"]


def indicadores_da_aba(ws):
    """-> {n: descrição} lendo os rótulos 'Indicador N:' do cabeçalho da aba de Registro"""
    out = {}
    for r in range(7, 21):
        rotulo = str(ws.cell(row=r, column=1).value or "").strip()
        if rotulo.startswith("Indicador"):
            n = re.sub(r"[^0-9]", "", rotulo)
            if n:
                out[n] = limpa(ws.cell(row=r, column=4).value)
    return out


# ---------------------------------------------------------------------------
# frases.json
# ---------------------------------------------------------------------------

def carregar_frases():
    if FRASES_JSON.exists():
        return json.loads(FRASES_JSON.read_text(encoding="utf-8"))
    return {}


def cmd_frases(trimestre=2):
    wb = abrir_wb()
    frases = carregar_frases()
    novas, revisar = 0, []
    for tab in sorted(TAB2UC, key=lambda t: int(t[2:])):
        reg, _, _ = abas_da_uc(wb, tab, trimestre)
        for n, desc in sorted(indicadores_da_aba(wb[reg]).items(), key=lambda kv: int(kv[0])):
            chave = f"{tab}/{n}"
            if chave in frases:
                continue
            v = flexionar(desc)
            frases[chave] = v
            novas += 1
            if v["PA"].startswith("??"):
                revisar.append(chave)
    FRASES_JSON.write_text(json.dumps(frases, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{FRASES_JSON}: {novas} entrada(s) nova(s), {len(frases)} no total.")
    print("Entradas já existentes foram preservadas — edições à mão não são sobrescritas.")
    if revisar:
        print("\nNão reconheci o verbo (marcadas com '??', corrija à mão):")
        for c in revisar:
            print(f"  {c}")


# ---------------------------------------------------------------------------
# geração / conferência
# ---------------------------------------------------------------------------

def montar(tab, trimestre=2):
    """-> dict com tudo que o .ts precisa, já validado"""
    wb = abrir_wb()
    reg, ind, proc = abas_da_uc(wb, tab, trimestre)
    uc_csv, uc_nome = TAB2UC[tab]
    descricoes = indicadores_da_aba(wb[reg])
    frases_todas = carregar_frases()

    frases = {}
    for n in descricoes:
        chave = f"{tab}/{n}"
        if chave not in frases_todas:
            sys.exit(f"falta '{chave}' em frases.json — rode:  python gerar.py frases")
        frases[n] = frases_todas[chave]
        for men in ("A", "PA", "NA", "NAV"):
            if frases[n][men].startswith("??"):
                sys.exit(f"'{chave}' menção {men} ainda está marcada '??' — revise frases.json")

    ordem, mencoes_csv = ler_csv()

    # erro de digitação no CSV vira abort aqui, não célula estranha na nuvem
    invalidos = []
    for nome_csv in ordem:
        for n in descricoes:
            v = mencoes_csv[nome_csv].get((uc_csv, f"{int(uc_csv[2:])}.{n}"))
            if v and v not in MENCOES_VALIDAS:
                invalidos.append(f"{uc_csv} col {int(uc_csv[2:])}.{n} — {nome_csv}: {v!r}")
    if invalidos:
        sys.exit("menções inválidas no CSV (esperado A/PA/NA/NAV):\n  " + "\n  ".join(invalidos))

    # confere a lista de alunos contra a aba de Registro
    ws = wb[reg]
    na_planilha = []
    for a, b in ((23, 27), (31, 42), (48, 60)):
        for r in range(a, b + 1):
            v = str(ws.cell(row=r, column=1).value or "").strip()
            if v:
                na_planilha.append(v)
    if [norm(x) for x in na_planilha] != [norm(x) for x in ordem]:
        divergencias = [(i + 1, a, b) for i, (a, b) in enumerate(zip(na_planilha, ordem)) if norm(a) != norm(b)]
        sys.exit(f"lista de alunos diverge entre CSV e '{reg}': {divergencias}")

    # a grafia da planilha é a canônica: é por ela que o .ts procura o aluno
    mencoes = {}
    for nome_csv, nome_pl in zip(ordem, na_planilha):
        m = {}
        for n in descricoes:
            if norm(nome_pl) == norm(EVADIDA):
                m[n] = "NAV"
                continue
            v = mencoes_csv[nome_csv].get((uc_csv, f"{int(uc_csv[2:])}.{n}"))
            if v:
                m[n] = v
        mencoes[nome_pl] = m
    ordem = na_planilha

    return {
        "tab": tab, "uc_csv": uc_csv, "uc_nome": uc_nome,
        "reg": reg, "ind": ind, "proc": proc,
        "descricoes": descricoes, "frases": frases,
        "ordem": ordem, "mencoes": mencoes, "wb": wb,
    }


def cmd_gerar(tab, dry, data_registro, trimestre):
    d = montar(tab, trimestre)
    if dry:
        return relatorio(d, data_registro)

    modelo = MODELO_TS.read_text(encoding="utf-8")
    subs = {
        "__TAB__": tab,
        "__UC_CSV__": d["uc_csv"],
        "__UC_NOME__": d["uc_nome"],
        "__DATA__": date.today().isoformat(),
        "__ABA_REGISTRO__": d["reg"],
        "__ABA_FB_IND__": d["ind"],
        "__ABA_FB_PROC__": d["proc"],
        "__PROC_PA__": PROC_PA,
        "__PROC_NA__": PROC_NA,
        "__DATA_REGISTRO__": data_registro,
        "__FRASES__": json.dumps(d["frases"], ensure_ascii=False, indent=4),
        "__MENCOES__": json.dumps(d["mencoes"], ensure_ascii=False, indent=4),
        "__SEM_PROCESSO__": json.dumps([EVADIDA], ensure_ascii=False),
    }
    for k, v in subs.items():
        modelo = modelo.replace(k, v)

    SAIDA.mkdir(exist_ok=True)
    destino = SAIDA / (f"preencher-{tab}.ts" if trimestre == 2 else f"preencher-{tab}-T{trimestre}.ts")
    destino.write_text(modelo, encoding="utf-8")

    lancados = sorted({n for m in d["mencoes"].values() for n in m}, key=int)
    pendentes = sorted(set(d["descricoes"]) - set(lancados), key=int)
    print(f"{destino}")
    print(f"  aba {tab} = {d['uc_csv']} ({d['uc_nome']})")
    print(f"  indicadores lançados: {', '.join(lancados) or '(nenhum)'}")
    if pendentes:
        print(f"  ainda sem menção no CSV, serão ignorados: {', '.join(pendentes)}")
    print("\nCole o conteúdo no Excel Online (Automatizar > Novo Script) e execute.")


def relatorio(d, data_registro):
    import openpyxl
    wb, tab = d["wb"], d["tab"]
    print(f"=== CONFERÊNCIA — aba {tab} = {d['uc_csv']} ({d['uc_nome']}) ===")
    print("nada é gravado; isto é o que o script escreveria na nuvem\n")

    lancados = sorted({n for m in d["mencoes"].values() for n in m}, key=int)
    pendentes = sorted(set(d["descricoes"]) - set(lancados), key=int)
    if pendentes:
        print(f"indicadores ainda sem menção no CSV (ignorados): {', '.join(pendentes)}\n")

    print(f"--- {d['reg']} ---")
    for n in lancados:
        dist = {}
        for nome in d["ordem"]:
            v = d["mencoes"][nome].get(n)
            if v:
                dist[v] = dist.get(v, 0) + 1
        col = 8 + 3 * sorted(d["descricoes"], key=int).index(n)
        letra_p = openpyxl.utils.get_column_letter(col)
        letra_f = openpyxl.utils.get_column_letter(col + 2)
        total = sum(dist.values()) * 2
        print(f"  Indicador {n}: {total} células ({letra_p}=P, {letra_f}=F)  "
              + ", ".join(f"{k}={v}" for k, v in sorted(dist.items(), key=lambda x: -x[1])))

    print(f"\n--- {d['ind']} ---")
    ws = wb[d["ind"]]
    for n in lancados:
        preenchidas = sum(1 for nome in d["ordem"] if d["mencoes"][nome].get(n))
        print(f"  Indicador {n}: {preenchidas} células")
        for men in ("A", "PA", "NA", "NAV"):
            usados = sum(1 for nome in d["ordem"] if d["mencoes"][nome].get(n) == men)
            if usados:
                print(f"     {men} ({usados}x): {d['frases'][n][men]}")

    print(f"\n--- {d['proc']} ---")
    ws = wb[d["proc"]]
    slots = [r for r in range(1, ws.max_row + 1)
             if str(ws.cell(row=r, column=1).value or "").strip() == "Nome:"]
    ultimo = max((i for i, r in enumerate(slots)
                  if str(ws.cell(row=r, column=3).value or "").strip()), default=-1)
    inicio = ultimo + 2
    pend = []
    for nome in d["ordem"]:
        if norm(nome) == norm(EVADIDA):
            continue
        vals = list(d["mencoes"][nome].values())
        if "NA" in vals:
            pend.append((nome, "NA"))
        elif "PA" in vals:
            pend.append((nome, "PA"))
    print(f"  último nome do T1: linha {slots[ultimo] if ultimo >= 0 else '(nenhum)'}")
    print(f"  T2 começa na linha {slots[inicio]}  ({len(pend)} linhas; data {data_registro} na coluna AB)")
    for i, (nome, men) in enumerate(pend):
        print(f"    r{slots[inicio + i]:<4} {men:<3} {nome}")
    semp = [n for n in d["ordem"] if norm(n) == norm(EVADIDA)]
    todoA = [n for n in d["ordem"]
             if norm(n) != norm(EVADIDA)
             and d["mencoes"][n]
             and all(v == "A" for v in d["mencoes"][n].values())]
    print(f"  sem linha: {len(todoA)} com A em tudo, {len(semp)} evadida(s)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("alvo", help="aba (UC1..UC9) ou 'frases'")
    ap.add_argument("--dry-run", action="store_true", help="só relata, não gera o .ts")
    ap.add_argument("--data", default=date.today().strftime("%d/%m/%Y"),
                    help="data do registro no Feedback de Processo (padrão: hoje)")
    ap.add_argument("--trimestre", type=int, default=2, choices=(1, 2, 3),
                    help="qual aba de Registro de Avaliação usar (padrão: 2)")
    args = ap.parse_args()

    for p in (CSV_PATH, XLSX_PATH):
        if not p.exists():
            sys.exit(f"não encontrei: {p}")

    if args.alvo == "frases":
        return cmd_frases(args.trimestre)
    if args.alvo not in TAB2UC:
        sys.exit(f"alvo inválido: {args.alvo}. Use 'frases' ou uma de: {', '.join(sorted(TAB2UC))}")
    cmd_gerar(args.alvo, args.dry_run, args.data, args.trimestre)


if __name__ == "__main__":
    main()
